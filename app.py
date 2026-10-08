import io
import re
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
import pandas as pd
import pypdf
import streamlit as st

st.set_page_config(page_title="Abertura de Conta Salario - Mirantes", page_icon="🏢", layout="wide")

# --- 1. TELA DE LOGIN ---
if "autenticado" not in st.session_state:
    st.session_state.autenticado = False

if not st.session_state.autenticado:
    col1, col2, col3 = st.columns([1, 2, 1])
    with col2:
        try:
            st.image("logo_mirantes.png", width=180)
        except Exception:
            st.html("<div>🏢 MIRANTES EMPREENDIMENTOS</div>")
        st.write("")
        st.html("<div>🔒 Acesso Restrito ao Departamento Pessoal</div>")
        st.text_input("Usuario:", value="dp@soumirantes.com.br", disabled=True)
        senha_input = st.text_input("Senha:", type="password", placeholder="••••••••")
        if st.button("Entrar no Sistema", use_container_width=True):
            if senha_input == "dp@mirantes":
                st.session_state.autenticado = True
                st.session_state.usuario = "dp@soumirantes.com.br"
                st.rerun()
            else:
                st.error("❌ Senha incorreta!")
    st.stop()

# --- 2. ÁREA LOGADA DA APLICAÇÃO ---
col_logo, col_titulo, col_user = st.columns([1, 3, 1])
with col_logo:
    try:
        st.image("logo_mirantes.png", width=140)
    except Exception:
        st.html("<div>🏢 MIRANTES</div>")

with col_titulo:
    st.title("Extracao de Dossies - Abertura de Conta Salario")

with col_user:
    st.caption("👤 **DP Mirantes**")
    st.caption(f"✉ {st.session_state.usuario}")
    if st.button("Sair", key="logout_btn"):
        st.session_state.autenticado = False
        st.rerun()

st.divider()

def extrair_dados_pdf(pdf_file, file_name):
    reader = pypdf.PdfReader(pdf_file)
    texto_completo = ""
    for page in reader.pages:
        txt = page.extract_text()
        if txt:
            texto_completo += "\n" + txt

    # 1. EXTRAÇÃO DINÂMICA DO NOME
    nome = ""
    lines = [line.strip() for line in texto_completo.split("\n") if line.strip()]
    for i, line in enumerate(lines):
        if line.upper() == "NOME" and i + 1 < len(lines):
            candidato = lines[i + 1].strip().upper()
            if candidato not in ["CPF", "MATRICULA", "TIPO DE REGISTRO", "ADMISSAO"] and len(candidato) > 3:
                nome = candidato
                break

    if not nome:
        nome_match = re.search(r"Nome civil[\s\n]+([A-ZÁÀÂÃÉÈÊÍÏÓÔÕÖÚÇ\s]{5,50})", texto_completo, re.IGNORECASE)
        if nome_match and "CPF" not in nome_match.group(1).upper():
            nome = nome_match.group(1).strip().split("\n")[0].upper()
        else:
            nome_limpo = re.sub(r"^\d+\s*", "", file_name.replace(".pdf", ""))
            nome = nome_limpo.upper()

    tx_up = texto_completo.upper()

    # 2. EXTRAÇÃO DINÂMICA DO CPF
    cpf_match = re.search(r"CPF[\s\n]*(\d{3}\.\d{3}\.\d{3}-\d{2})", texto_completo) or re.search(r"\b(\d{3}\.\d{3}\.\d{3}-\d{2})\b", texto_completo)
    if cpf_match:
        cpf_num = int(re.sub(r"\D", "", cpf_match.group(1)))
    else:
        cpf_digits = re.search(r"\b(\d{11})\b", texto_completo)
        cpf_num = int(cpf_digits.group(1)) if cpf_digits else ""

    # 3. EXTRAÇÃO DINÂMICA DO TELEFONE & DDD
    ddd = 84
    tel_match = re.search(r"(9\d{4}[-\s]?\d{4})", texto_completo)
    if tel_match and "90363" not in tel_match.group(1) and "90027" not in tel_match.group(1):
        tel_raw = re.sub(r"\D", "", tel_match.group(1))
        telefone = f"{tel_raw[:5]}-{tel_raw[5:]}"
    else:
        telefone = ""

    # 4. EXTRAÇÃO DINÂMICA DO E-MAIL
    email_match = re.search(r"([a-zA-Z0-9._%+-]+@(gmail|hotmail|outlook|yahoo|live|icloud)[a-zA-Z0-9.-]*\.[a-zA-Z]{2,})", texto_completo, re.IGNORECASE)
    email = email_match.group(1).lower() if email_match else ""

    # 5. EXTRAÇÃO DINÂMICA DO RG / CIN (ÓRGÃO EXPEDITO FIXO EM SSP)
    orgao_expeditor = "SSP"
    doc_num_match = re.search(r"REGISTRO GERAL[\s\n]*([\d.]+)", texto_completo, re.IGNORECASE)
    if doc_num_match and "000000" not in doc_num_match.group(0):
        doc_numero = re.sub(r"\D", "", doc_num_match.group(0)).zfill(9)
    else:
        doc_numero = str(cpf_num).zfill(11) if cpf_num else ""
    
    # Busca data de emissão no texto se houver
    dt_match = re.search(r"\b(\d{2}/\d{2}/\d{4})\b", texto_completo)
    data_emissao = dt_match.group(1) if dt_match else ""

    # 6. DADOS BANCÁRIOS E REGRAS AUTOMÁTICAS POR BANCO
    banco = ""
    prod_operacao = ""
    agencia = ""
    conta = ""
    dv = ""

    if "SANTANDER" in tx_up or "033" in tx_up:
        banco = 33
        prod_operacao = ""
    elif "BANCO DO BRASIL" in tx_up or "001" in tx_up or " 001 " in tx_up:
        banco = 1
        prod_operacao = ""
    elif "ITAU" in tx_up or "341" in tx_up:
        banco = 341
        prod_operacao = ""
    elif "CAIXA" in tx_up or "104" in tx_up:
        banco = 104
        if "POUPANCA" in tx_up or "POUPANÇA" in tx_up:
            prod_operacao = "CONTA POUPANÇA"
        elif "CONTA FACIL" in tx_up or "CONTA FÁCIL" in tx_up:
            prod_operacao = "CONTA FÁCIL"
        else:
            prod_operacao = "CONTA CORRENTE"

    # Busca agência e conta se existirem no texto do PDF
    ag_match = re.search(r"AGENCIA[\s:]*(\d{3,5})", tx_up)
    if ag_match:
        agencia = ag_match.group(1)

    cc_match = re.search(r"CONTA[\s:]*(\d{4,10})[-_]?(\d{1})?", tx_up)
    if cc_match:
        conta = cc_match.group(1)
        dv = cc_match.group(2) if cc_match.group(2) else ""

    return {
        "NOME": nome,
        "CPF": cpf_num,
        "DDD": ddd,
        "TELEFONE": telefone,
        "E-MAIL": email,
        "DOC - NÚMERO": doc_numero,
        "DOC - ÓRGÃO EXPEDITO": orgao_expeditor,
        "DOC - UF ÓRGÃO EMISSOR": "RN",
        "DOC - DATA DE EMISSÃO": data_emissao,
        "DOC - DATA DE VALIDADE": "",
        "CONTA SALÁRIO - AGÊNCIA": "",
        "CONTA SALÁRIO - PROD/OPERAÇÃO": "",
        "CONTA SALÁRIO - CONTA": "",
        "CONTA SALÁRIO - DV": "",
        "CONTA DESTINO - BANCO": banco,
        "CONTA DESTINO - AGÊNCIA": agencia,
        "CONTA DESTINO - PROD/OPERAÇÃO": prod_operacao,
        "CONTA DESTINO - CONTA": conta,
        "CONTA DESTINO - DV": dv,
    }

uploaded_files = st.file_uploader("Selecione um ou varios PDFs de funcionarios de uma vez", type=["pdf"], accept_multiple_files=True)

if uploaded_files:
    registros = []
    for file in uploaded_files:
        dados = extrair_dados_pdf(file, file.name)
        registros.append(dados)

    df = pd.DataFrame(registros)
    st.subheader("Pre-visualizacao dos Dados Extraidos")
    st.dataframe(df, use_container_width=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Abertura de Conta"

    font_bold = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold_red = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_regular = Font(name="Calibri", size=11, color="000000")

    fill_red = PatternFill(start_color="FF3300", end_color="FF3300", fill_type="solid")
    fill_blue = PatternFill(start_color="0000CC", end_color="0000CC", fill_type="solid")

    border_thin = Border(left=Side(style="thin", color="D9D9D9"), right=Side(style="thin", color="D9D9D9"), top=Side(style="thin", color="D9D9D9"), bottom=Side(style="thin", color="D9D9D9"))

    ws["A1"] = "CNPJ"
    ws["A1"].font = font_bold_red
    ws["A1"].fill = fill_red
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws["B1"] = "49.036.333/0001-60"
    ws["B1"].font = font_regular
    ws["B1"].alignment = Alignment(horizontal="left", vertical="center")

    headers = ["NOME", "CPF", "DDD", "TELEFONE", "E-MAIL", "DOC - NÚMERO", "DOC - ÓRGÃO EXPEDITO", "DOC - UF ÓRGÃO EMISSOR", "DOC - DATA DE EMISSÃO", "DOC - DATA DE VALIDADE", "CONTA SALÁRIO - AGÊNCIA", "CONTA SALÁRIO - PROD/OPERAÇÃO", "CONTA SALÁRIO - CONTA", "CONTA SALÁRIO - DV", "CONTA DESTINO - BANCO", "CONTA DESTINO - AGÊNCIA", "CONTA DESTINO - PROD/OPERAÇÃO", "CONTA DESTINO - CONTA", "CONTA DESTINO - DV"]

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_num)
        cell.value = header
        cell.font = font_bold
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = fill_red if col_num <= 2 else fill_blue

    # Lista Suspensa na Coluna Q (PROD/OPERAÇÃO)
    dv_operacao = DataValidation(type="list", formula1='"CONTA CORRENTE,CONTA POUPANÇA,CONTA FÁCIL"', allow_blank=True)
    ws.add_data_validation(dv_operacao)
    dv_operacao.add("Q3:Q200")

    fmt_cpf = '000"."000"."000"-"00'

    for row_idx, row_data in enumerate(registros, start=3):
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = row_data.get(header, "")
            cell.font = font_regular
            cell.border = border_thin
            cell.alignment = Alignment(horizontal="left", vertical="center")

            if header == "CPF" and val != "":
                try:
                    cell.value = int(str(val).replace("-", "").replace(".", ""))
                    cell.number_format = fmt_cpf
                except Exception:
                    cell.value = str(val)
            else:
                cell.value = val

            if header == "DOC - NÚMERO":
                cell.number_format = "@"

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    excel_data = output.getvalue()

    st.download_button(label="📥 Baixar Planilha Padrao Caixa Economica (.xlsx)", data=excel_data, file_name="Planilha_Abertura_Conta_Salario_Caixa.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
